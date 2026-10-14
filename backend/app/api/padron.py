"""Padrón de alumnos: alta, carga masiva y depuración por antigüedad.

El padrón es lo que cierra el sistema: quien no está aquí no puede activar
cuenta, así que no hay registro abierto ni alumnos de otras facultades.
"""

import csv
import io
from datetime import date, datetime
from pathlib import Path

from flask import Blueprint, jsonify, request, current_app
from flask_login import current_user

from ..extensions import db
from ..models import PadronAlumno, Usuario, Rol, Tramite, Archivo, Evento, TipoEvento
from ..seguridad import rol_requerido
from ..dominio.servicios import ErrorDominio
from ..almacenamiento import borrar

bp = Blueprint("padron", __name__)

COLUMNAS = ["matricula", "curp_prefijo", "nombre", "correo_institucional"]
OPCIONALES = ["fecha_ingreso"]


@bp.errorhandler(ErrorDominio)
def _error(e):
    db.session.rollback()
    return jsonify({"error": str(e)}), 400


def _anios_vigencia():
    return current_user.programa.anios_vigencia if current_user.programa else 6


def _antiguedad(registro):
    hoy = date.today()
    inicio = registro.inicio
    return (hoy - inicio).days / 365.25


def _tramites_de(matricula):
    usuario = Usuario.query.filter_by(identificador=matricula, rol=Rol.ALUMNO).first()
    return Tramite.query.filter_by(alumno_id=usuario.id).count() if usuario else 0


def _serializar(r):
    anios = _antiguedad(r)
    tramites = _tramites_de(r.matricula)
    return {
        "id": r.id, "matricula": r.matricula, "nombre": r.nombre,
        "correo": r.correo_institucional, "activado": r.activado,
        "dado_de_baja": r.dado_de_baja,
        "fecha_ingreso": r.inicio.isoformat(),
        "anios": round(anios, 1),
        "vencido": anios >= _anios_vigencia(),
        "tramites": tramites,
        # Solo quien no dejó expediente. La pantalla no decide esto por su
        # cuenta: la regla vive en el servidor, que es quien la aplica.
        "eliminable": tramites == 0,
    }


@bp.get("")
@rol_requerido(Rol.SECRETARIA, Rol.COORDINADORA)
def listar():
    consulta = PadronAlumno.query
    if current_user.programa_id:
        consulta = consulta.filter_by(programa_id=current_user.programa_id)
    registros = consulta.order_by(PadronAlumno.matricula).all()
    datos = [_serializar(r) for r in registros]
    return jsonify({
        "total": len(datos),
        "activados": sum(1 for d in datos if d["activado"]),
        "vencidos": sum(1 for d in datos if d["vencido"] and not d["dado_de_baja"]),
        "anios_vigencia": _anios_vigencia(),
        "alumnos": datos,
    })


# --------------------------------------------------------------------------
# Alta individual
# --------------------------------------------------------------------------

def _leer_fecha(valor):
    if not valor:
        return None
    if isinstance(valor, (date, datetime)):
        return valor.date() if isinstance(valor, datetime) else valor
    texto = str(valor).strip()
    if not texto:
        return None
    # Solo el año también sirve: en la práctica es lo que suele haber a mano.
    if texto.isdigit() and len(texto) == 4:
        return date(int(texto), 8, 1)
    for formato in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            continue
    raise ErrorDominio(f"No entiendo la fecha de ingreso «{texto}». Usa AAAA-MM-DD o solo el año.")


def _validar(matricula, curp, nombre, correo):
    if not matricula or not nombre or not correo:
        raise ErrorDominio("Matrícula, nombre y correo son obligatorios.")
    if len(curp) != 10:
        raise ErrorDominio("El prefijo de CURP debe tener exactamente 10 caracteres.")
    if "@" not in correo:
        raise ErrorDominio("El correo institucional no es válido.")


@bp.post("")
# Coordinación también da de alta: en la práctica registra alumnos igual que
# Secretaría, y mandarla a pedirlo era un rodeo sin motivo.
@rol_requerido(Rol.SECRETARIA, Rol.COORDINADORA)
def agregar():
    """Alta de uno en uno: no siempre hay un archivo, y para un alumno que
    llega tarde o al que se le corrigió la matrícula, preparar una hoja de
    cálculo es más trabajo que el alta misma."""
    datos = request.get_json() or {}
    matricula = (datos.get("matricula") or "").strip().upper()
    curp = (datos.get("curp_prefijo") or "").strip().upper()
    nombre = (datos.get("nombre") or "").strip()
    correo = (datos.get("correo_institucional") or "").strip()
    _validar(matricula, curp, nombre, correo)

    if PadronAlumno.query.filter_by(matricula=matricula).first():
        raise ErrorDominio(f"La matrícula {matricula} ya está en el padrón.")

    registro = PadronAlumno(
        matricula=matricula, curp_prefijo=curp, nombre=nombre,
        correo_institucional=correo, programa_id=current_user.programa_id,
        fecha_ingreso=_leer_fecha(datos.get("fecha_ingreso")),
    )
    db.session.add(registro)
    db.session.commit()
    return jsonify(_serializar(registro)), 201


# --------------------------------------------------------------------------
# Carga masiva
# --------------------------------------------------------------------------

def _filas_csv(contenido):
    try:
        texto = contenido.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise ErrorDominio(
            "El archivo no está en UTF-8. Si lo exportaste de Excel, guárdalo "
            "como «CSV UTF-8» para que los acentos no se corrompan, o sube "
            "directamente el .xlsx."
        )
    lector = csv.DictReader(io.StringIO(texto))
    return (lector.fieldnames or []), list(lector)


def _filas_excel(contenido):
    """Excel evita el problema de codificación del CSV: los acentos y la ñ
    viajan bien sin que nadie tenga que elegir el formato correcto al exportar."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ErrorDominio("Falta la biblioteca openpyxl. Instálala con pip.")

    hoja = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True).active
    renglones = hoja.iter_rows(values_only=True)
    try:
        encabezados = [str(c).strip().lower() if c is not None else "" for c in next(renglones)]
    except StopIteration:
        raise ErrorDominio("La hoja está vacía.")

    filas = []
    for valores in renglones:
        if all(v is None or str(v).strip() == "" for v in valores):
            continue
        filas.append(dict(zip(encabezados, valores)))
    return encabezados, filas


@bp.post("/cargar")
@rol_requerido(Rol.SECRETARIA, Rol.COORDINADORA)
def cargar():
    """Acepta CSV y Excel con: matricula, curp_prefijo, nombre,
    correo_institucional, y opcionalmente fecha_ingreso.

    Procesa fila por fila y devuelve el detalle de lo que falló, en lugar de
    rechazar el archivo completo por un error: con cientos de renglones,
    «el archivo tiene un error» no sirve de nada.
    """
    archivo = request.files.get("archivo")
    if not archivo:
        raise ErrorDominio("Adjunta el archivo.")

    contenido = archivo.read()
    extension = Path(archivo.filename or "").suffix.lower()
    if extension in (".xlsx", ".xlsm"):
        encabezados, filas = _filas_excel(contenido)
    elif extension == ".csv":
        encabezados, filas = _filas_csv(contenido)
    else:
        raise ErrorDominio("Solo se aceptan archivos .csv o .xlsx.")

    faltantes = [c for c in COLUMNAS if c not in encabezados]
    if faltantes:
        raise ErrorDominio(f"Faltan columnas en el archivo: {', '.join(faltantes)}")

    creados, actualizados, errores = 0, 0, []

    for n, fila in enumerate(filas, start=2):
        def val(clave):
            v = fila.get(clave)
            return "" if v is None else str(v).strip()

        matricula, curp = val("matricula").upper(), val("curp_prefijo").upper()
        nombre, correo = val("nombre"), val("correo_institucional")

        try:
            _validar(matricula, curp, nombre, correo)
            ingreso = _leer_fecha(fila.get("fecha_ingreso"))
        except ErrorDominio as e:
            errores.append({"fila": n, "motivo": f"{e} ({matricula or 'sin matrícula'})"})
            continue

        registro = PadronAlumno.query.filter_by(matricula=matricula).first()
        if registro:
            # Nunca se toca `activado`: si el alumno ya definió su contraseña,
            # recargar el padrón no debe desactivarle la cuenta.
            registro.curp_prefijo, registro.nombre = curp, nombre
            registro.correo_institucional = correo
            if ingreso:
                registro.fecha_ingreso = ingreso
            actualizados += 1
        else:
            db.session.add(PadronAlumno(
                matricula=matricula, curp_prefijo=curp, nombre=nombre,
                correo_institucional=correo, fecha_ingreso=ingreso,
                programa_id=current_user.programa_id,
            ))
            creados += 1

    db.session.commit()
    return jsonify({"creados": creados, "actualizados": actualizados, "errores": errores})


# --------------------------------------------------------------------------
# Depuración por antigüedad
# --------------------------------------------------------------------------

def _vencidos():
    consulta = PadronAlumno.query.filter_by(dado_de_baja=False)
    if current_user.programa_id:
        consulta = consulta.filter_by(programa_id=current_user.programa_id)
    limite = _anios_vigencia()
    return [r for r in consulta.all() if _antiguedad(r) >= limite]


@bp.delete("/<int:registro_id>")
@rol_requerido(Rol.SECRETARIA, Rol.COORDINADORA)
def eliminar_alumno(registro_id):
    """Quita a un alumno del padrón, con la misma regla que la depuración:
    solo si no hizo trámites.

    Para el alta equivocada —una matrícula mal escrita, un alumno que nunca
    perteneció al programa— y para limpiar los datos de prueba antes de abrir
    el sistema. Si ya hizo trámites, su expediente es contabilidad de la
    facultad y lo que corresponde es darlo de baja, no borrarlo.
    """
    registro = db.session.get(PadronAlumno, registro_id)
    if not registro or (current_user.programa_id
                        and registro.programa_id != current_user.programa_id):
        return jsonify({"error": "Alumno no encontrado"}), 404

    usuario = Usuario.query.filter_by(identificador=registro.matricula,
                                      rol=Rol.ALUMNO).first()
    tramites = Tramite.query.filter_by(alumno_id=usuario.id).count() if usuario else 0
    if tramites:
        raise ErrorDominio(
            f"{registro.nombre} tiene {tramites} trámite(s) registrados. "
            "Eliminarlo dejaría folios sin titular: da de baja su cuenta, que "
            "le quita el acceso y conserva el expediente."
        )

    nombre = registro.nombre
    if usuario:
        db.session.delete(usuario)
    db.session.delete(registro)
    db.session.commit()
    return jsonify({"eliminado": nombre})


@bp.get("/depurar")
@rol_requerido(Rol.COORDINADORA)
def revisar_depuracion():
    """Vista previa antes de ejecutar. Borrar datos de alumnos no es algo que
    deba pasar sin que alguien vea antes exactamente a quiénes alcanza."""
    salida = []
    for r in _vencidos():
        usuario = Usuario.query.filter_by(identificador=r.matricula, rol=Rol.ALUMNO).first()
        tramites = Tramite.query.filter_by(alumno_id=usuario.id).count() if usuario else 0
        archivos = (db.session.query(Archivo).join(Tramite)
                    .filter(Tramite.alumno_id == usuario.id).count()) if usuario else 0
        salida.append({
            **_serializar(r),
            "tramites": tramites, "archivos": archivos,
            # Quien nunca usó el sistema no deja rastro que conservar.
            "accion": "Se elimina por completo" if tramites == 0
                      else "Pierde el acceso y se borran sus archivos",
        })
    return jsonify({"anios_vigencia": _anios_vigencia(), "alumnos": salida})


@bp.post("/depurar")
@rol_requerido(Rol.COORDINADORA)
def depurar():
    """Dos destinos distintos, y la diferencia importa:

    - Quien nunca hizo un trámite se elimina por completo. No hay nada que
      conservar y su nombre y correo no tienen por qué seguir ahí.
    - Quien sí hizo trámites pierde el acceso y se borran SUS ARCHIVOS
      —comprobantes con datos bancarios, documentos escaneados—, pero el
      registro del trámite se queda: folio, tipo, fechas y quién lo atendió son
      contabilidad de la facultad, y borrarlos dejaría huecos imposibles de
      aclarar después.
    """
    if (request.get_json() or {}).get("confirmar") is not True:
        raise ErrorDominio("Falta confirmar la depuración.")

    eliminados, dados_de_baja, archivos_borrados = 0, 0, 0

    for registro in _vencidos():
        usuario = Usuario.query.filter_by(identificador=registro.matricula, rol=Rol.ALUMNO).first()
        tramites = Tramite.query.filter_by(alumno_id=usuario.id).all() if usuario else []

        if not tramites:
            if usuario:
                db.session.delete(usuario)
            db.session.delete(registro)
            eliminados += 1
            continue

        for tramite in tramites:
            for archivo in list(tramite.archivos):
                borrar(archivo.ruta_relativa)
                db.session.delete(archivo)
                archivos_borrados += 1
            # Queda constancia de la depuración en la propia bitácora.
            db.session.add(Evento(
                tramite_id=tramite.id, tipo=TipoEvento.COMENTARIO,
                texto="Archivos depurados por antigüedad del alumno.",
                autor_id=current_user.id,
            ))

        registro.dado_de_baja = True
        usuario.activo = False
        dados_de_baja += 1

    db.session.commit()
    return jsonify({
        "eliminados": eliminados,
        "dados_de_baja": dados_de_baja,
        "archivos_borrados": archivos_borrados,
    })
